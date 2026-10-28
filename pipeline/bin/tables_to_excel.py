#!/usr/bin/env python3
"""tables_to_excel.py out.xlsx "Sheet name=table.tsv" ...  (missing tables are skipped)
Also concatenates per-species tables given as "Sheet=glob*.tsv"."""
import csv, glob, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

out = sys.argv[1]
wb = Workbook(); wb.remove(wb.active)
for spec in sys.argv[2:]:
    name, pat = spec.split("=", 1)
    files = sorted(glob.glob(pat))
    rows, hdr = [], None
    for f in files:
        with open(f) as fh:
            r = list(csv.reader(fh, delimiter="\t"))
        if not r:
            continue
        if hdr is None:
            hdr = r[0]; rows += r[1:]
        else:
            rows += r[1:] if r[0] == hdr else r
    if hdr is None:
        print(f"  skip  {name:34s} (not produced yet)"); continue
    ws = wb.create_sheet(name[:31]); ws.append(hdr)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="2F4B7C")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        vals = []
        for v in r:
            try:
                vals.append(int(v) if v.lstrip("-").isdigit() else float(v))
            except ValueError:
                vals.append(v if len(v) < 32000 else v[:32000] + " ...[truncated]")
        ws.append(vals)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(max(10, max(len(str(c.value or "")) for c in col[:200]) + 2), 45)
    ws.freeze_panes = "B2"
    print(f"  ok    {name:34s} {len(rows):>7} rows")
wb.save(out)
print("Excel:", out)
